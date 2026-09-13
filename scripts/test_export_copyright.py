"""Regression tests for the source-material exporter; see docs/copyright-materials.md."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

import export_copyright as exporter


class CopyrightExportTests(unittest.TestCase):
    def test_source_scope(self):
        for path in ["crates/core/src/lib.rs", "apps/desktop/src/新增.vue"]:
            self.assertTrue(exporter.is_source(path))
        for path in ["third_party/src/a.rs", "scripts/a.py",
                     "apps/desktop/src/a.test.ts", "crates/core/src/tests.rs",
                     "apps/desktop/src/__tests__/a.ts", "apps/desktop/src/a.png",
                     "apps/desktop/node_modules/pkg/src/a.ts"]:
            self.assertFalse(exporter.is_source(path), path)

    def test_counts_order_bom_and_original_line_numbers(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder).resolve()
            names = ["apps/desktop/src/新增.vue", "crates/core/src/lib.rs"]
            for name in names:
                target = root / name
                target.parent.mkdir(parents=True)
                target.write_bytes("\ufeff// comment\r\n\r\n  code\r\n".encode("utf-8"))
            with patch.object(exporter, "git", return_value="\0".join(names) + "\0"):
                files, rows = exporter.collect_sources(root)
            self.assertEqual([file["file"] for file in files], names[::-1])
            self.assertEqual(sum(file["physical_lines"] for file in files), 6)
            self.assertEqual(len(rows), 4)
            self.assertEqual(rows[1], (names[1], 3, "  code"))

    def test_untracked_is_explicit(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch.object(exporter, "git", return_value="") as git:
                with self.assertRaises(ValueError):
                    exporter.collect_sources(Path(folder), True)
                self.assertIn("--others", git.call_args.args)
                self.assertIn("--exclude-standard", git.call_args.args)

    def test_page_boundaries_and_no_overlap(self):
        for count in (1, 49, 50, 51, 2999, 3000, 3001, 8000):
            rows = [("a.rs", number, "code") for number in range(1, count + 1)]
            pages = exporter.select_pages(rows)
            selected = [row for _, _, block in pages for row in block]
            expected = rows if count <= 3000 else rows[:1500] + rows[-1500:]
            self.assertEqual(selected, expected)
            self.assertEqual(len(set(row[1] for row in selected)), len(selected))
            self.assertTrue(all(len(block) == 50 for _, _, block in pages[:-1]))

    def test_docx_text_and_explicit_pages(self):
        rows = [("a.rs", number, "\tlet name = \"中文\";") for number in range(1, 3002)]
        pages = exporter.select_pages(rows)
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "source.docx"
            exporter.export_docx(output, pages, "DocVault", "9.8.7")
            doc = Document(output)
            self.assertEqual(sum(bool(p.paragraph_format.page_break_before)
                                 for p in doc.paragraphs), 59)
            code = [p.text for p in doc.paragraphs if p.text[:6].isdigit()]
            self.assertEqual(len(code), 3000)
            self.assertEqual(code[0], '000001      let name = "中文";')
            self.assertEqual(code[-1], '003001      let name = "中文";')
            self.assertIn("9.8.7", doc.sections[0].header.paragraphs[0].text)

    def test_extreme_lines_fail_instead_of_truncating(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "source.docx"
            with self.assertRaisesRegex(ValueError, "a.rs:1"):
                exporter.export_docx(path, [("全部", 0, [("a.rs", 1, "x" * 2000)])],
                                     "DocVault", "1")
            self.assertFalse(path.exists())

    def test_existing_output_is_untouched(self):
        with tempfile.TemporaryDirectory() as folder:
            sentinel = Path(folder) / "registration.docx"
            sentinel.write_bytes(b"original")
            self.assertEqual(exporter.main(["--output-dir", folder]), 1)
            self.assertEqual(sentinel.read_bytes(), b"original")

    def test_main_auto_version_statistics_and_failure_log(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            package = root / "apps/desktop/package.json"
            package.parent.mkdir(parents=True)
            package.write_text('{"version":"2.3.4"}', encoding="utf-8")
            source = root / "crates/core/src/lib.rs"
            source.parent.mkdir(parents=True)
            source.write_text("// header\n\nfn main() {}\n", encoding="utf-8")

            def git(_, *args):
                return {"ls-files": "crates/core/src/lib.rs\0", "rev-parse": "abc123\n",
                        "status": " M crates/core/src/lib.rs\n"}[args[0]]

            output = root / "result"
            with patch.object(exporter, "git", side_effect=git):
                self.assertEqual(exporter.main(["--root", str(root), "--output-dir",
                                                str(output), "--stats-only"]), 0)
            report = json.loads((output / "source-statistics.json").read_text(encoding="utf-8"))
            self.assertEqual((report["version"], report["nonblank_lines"],
                              report["physical_lines"]), ("2.3.4", 2, 3))
            self.assertTrue(report["working_tree_dirty"])
            self.assertEqual(list(output.glob("*.docx")), [])
            failed = root / "failed"
            with patch.object(exporter, "git", side_effect=RuntimeError("Git failure")):
                self.assertEqual(exporter.main(["--root", str(root), "--output-dir",
                                                str(failed)]), 1)
            self.assertIn("Git failure", (failed / "generation.log").read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()
