# /// script
# requires-python = ">=3.11"
# dependencies = ["python-docx==1.2.0"]
# ///
"""Count DocVault source and export traceable copyright program materials."""

import argparse
import csv
import hashlib
import json
import logging
import subprocess
import sys
from datetime import datetime
from pathlib import Path, PurePosixPath

SOURCE_ROOTS = (
    "shared/types/", "crates/core/", "crates/storage/", "crates/ooxml/",
    "crates/jobs/", "apps/cli/", "apps/desktop/src-tauri/",
    "apps/office-addin-cli/", "apps/office-addin/", "shared/addin-web/",
    "apps/desktop/src/",
)
EXTENSIONS = {".rs", ".ts", ".vue", ".css", ".js", ".html"}
LINES_PER_PAGE = 50
HALF_LINES = 1500


def git(root, *args):
    result = subprocess.run(
        ["git", "-C", str(root), *args], capture_output=True, check=False
    )
    if result.returncode:
        raise RuntimeError(result.stderr.decode("utf-8", errors="replace").strip())
    return result.stdout.decode("utf-8")


def is_source(name):
    path = PurePosixPath(name)
    return (
        path.parts[0] in {"apps", "crates", "shared"}
        and "src" in path.parts
        and path.suffix in EXTENSIONS
        and not set(path.parts).intersection(
            {"tests", "test", "__tests__", "node_modules", "target", "dist", "vendor"}
        )
        and ".test." not in path.name
        and ".spec." not in path.name
        and path.name != "tests.rs"
    )


def collect_sources(root, include_untracked=False):
    args = ["ls-files", "-z", "--cached"]
    if include_untracked:
        args.extend(["--others", "--exclude-standard"])
    names = {name for name in git(root, *args).split("\0") if name and is_source(name)}
    names = sorted(names, key=lambda name: (
        next((i for i, prefix in enumerate(SOURCE_ROOTS) if name.startswith(prefix)),
             len(SOURCE_ROOTS)), name,
    ))
    files, rows = [], []
    for name in names:
        path = root / name
        if path.is_symlink() or not path.resolve().is_relative_to(root):
            raise ValueError(f"Source must be a regular file inside repository: {name}")
        if not path.exists():
            logging.warning("Skipping tracked file deleted from working tree: %s", name)
            continue
        raw = path.read_bytes()
        lines = raw.decode("utf-8-sig").splitlines()
        selected = [(name, number, text) for number, text in enumerate(lines, 1)
                    if text.strip()]
        rows.extend(selected)
        files.append({"file": name, "physical_lines": len(lines),
                      "nonblank_lines": len(selected),
                      "sha256": hashlib.sha256(raw).hexdigest()})
    if not rows:
        raise ValueError("No nonblank source lines found in the selected repository.")
    return files, rows


def select_pages(rows):
    """No duplicated overlap for programs of 3000 lines or fewer."""
    if len(rows) <= HALF_LINES * 2:
        blocks = [("全部", 0, rows)]
    else:
        blocks = [("前段", 0, rows[:HALF_LINES]),
                  ("后段", len(rows) - HALF_LINES, rows[-HALF_LINES:])]
    return [(label, start + offset, block[offset:offset + LINES_PER_PAGE])
            for label, start, block in blocks
            for offset in range(0, len(block), LINES_PER_PAGE)]


def export_docx(path, pages, name, version, landscape=False):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def font(run, size, code=False):
        run.font.name = "Consolas" if code else "宋体"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.top_margin = section.bottom_margin = Cm(1.6)
    section.left_margin = section.right_margin = Cm(1.5)
    doc.core_properties.title = f"{name} {version} 源程序材料"
    doc.core_properties.author = ""
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(header.add_run(doc.core_properties.title), 9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("第 "), 9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    font(footer.add_run(f" 页／共 {len(pages)} 页"), 9)
    width = section.page_width.pt - section.left_margin.pt - section.right_margin.pt - 8
    for page_number, (label, start, rows) in enumerate(pages, 1):
        title = doc.add_paragraph()
        title.paragraph_format.page_break_before = page_number > 1
        title.paragraph_format.space_after = Pt(6)
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.line_spacing = Pt(12)
        font(title.add_run(f"{label}  源程序非空行 {start + 1}–{start + len(rows)}"), 9)
        for offset, (source, line_number, text) in enumerate(rows, start + 1):
            display = f"{offset:06d}  {text.expandtabs(4)}"
            # Conservative width estimate; never truncate long source lines.
            units = sum(0.65 if ord(char) < 256 else 1.1 for char in display)
            size = min(7.5, int(width / max(units, 1) * 2) / 2)
            if size < 5:
                raise ValueError(
                    f"Source line too wide to print legibly: {source}:{line_number}. "
                    "Try --landscape or format the source line before exporting."
                )
            paragraph = doc.add_paragraph()
            fmt = paragraph.paragraph_format
            fmt.space_after = fmt.space_before = Pt(0)
            fmt.line_spacing = 1.5
            fmt.keep_together = True
            fmt.widow_control = False
            font(paragraph.add_run(display), size, code=True)
    doc.save(path)


def write_index(path, pages):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["page", "page_line", "block", "source_nonblank_line",
                         "source_file", "source_line"])
        for page_number, (label, start, rows) in enumerate(pages, 1):
            for page_line, (source, source_line, _) in enumerate(rows, 1):
                writer.writerow([page_number, page_line, label, start + page_line,
                                 source, source_line])


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1],
                        help="Repository root (defaults to the script's repository)")
    parser.add_argument("--output-dir", type=Path, help="New output directory; must be empty")
    parser.add_argument("--name", default="兰天嗨彩办公文档管理（DocVault）")
    parser.add_argument("--version", help="Defaults to apps/desktop/package.json version")
    parser.add_argument("--include-untracked", action="store_true",
                        help="Include new source files not ignored by Git")
    parser.add_argument("--stats-only", action="store_true", help="Only produce statistics")
    parser.add_argument("--landscape", action="store_true", help="Use landscape A4 for long lines")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    root = args.root.resolve()
    stamp = datetime.now().astimezone()
    output = (args.output_dir or root / "dist" / "copyright-materials" /
              stamp.strftime("export-%Y%m%d-%H%M%S-%f")).resolve()
    # Never replace the registration document or an earlier source export.
    if output.exists() and (not output.is_dir() or any(output.iterdir())):
        print(f"Output directory must be empty: {output}", file=sys.stderr)
        return 1
    try:
        output.mkdir(parents=True, exist_ok=True)
        handler = logging.FileHandler(output / "generation.log", encoding="utf-8")
    except OSError as error:
        # No durable log is possible if the output/log directory is unwritable.
        print(f"Cannot create output or persistent log: {error}", file=sys.stderr)
        return 1
    handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
    logger = logging.getLogger()
    logger.addHandler(handler)
    logger.setLevel(logging.INFO)
    try:
        logging.info("Source export started: %s", root)
        version = args.version or json.loads(
            (root / "apps/desktop/package.json").read_text(encoding="utf-8-sig")
        )["version"]
        files, rows = collect_sources(root, args.include_untracked)
        pages = select_pages(rows)
        report = {
            "software_name": args.name, "version": version,
            "generated_at": stamp.isoformat(), "git_head": git(root, "rev-parse", "HEAD").strip(),
            "working_tree_dirty": bool(git(root, "status", "--porcelain", "--untracked-files=normal").strip()),
            "include_untracked": args.include_untracked,
            "counting_policy": "Nonblank lines include comments and inline tests; separate test files excluded.",
            "file_count": len(files), "physical_lines": sum(f["physical_lines"] for f in files),
            "nonblank_lines": len(rows), "planned_pages": len(pages),
            "selected_lines": sum(len(page[2]) for page in pages),
            "layout_visually_verified": False, "files": files,
        }
        (output / "source-statistics.json").write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        summary = (f"软件：{args.name}\n版本：{version}\n文件数：{len(files)}\n"
                   f"物理行：{report['physical_lines']}\n非空行：{len(rows)}\n"
                   "非空行含注释和文件内测试；不等于剔除注释后的有效代码行。\n"
                   f"材料行数：{report['selected_lines']}；计划页数：{len(pages)}。\n"
                   "大于 3000 行取前后各 1500 行，否则输出全部；每页 50 行，末页可不足。\n"
                   "读取当前工作区内容，文件顺序与 SHA-256 见 JSON，原始行号见 CSV。\n"
                   "Word 字体替换可能影响分页；导出后请检查实际页数及长行。\n")
        (output / "source-statistics.txt").write_text(summary, encoding="utf-8-sig")
        if not args.stats_only:
            export_docx(output / "DocVault-源程序材料.docx", pages, args.name, version, args.landscape)
            write_index(output / "source-line-index.csv", pages)
        logging.info("Source export completed: %s files, %s nonblank lines", len(files), len(rows))
        print(summary)
        print(f"Output: {output}")
        return 0
    except Exception:
        logging.exception("Source export failed")
        print(f"Export failed; details: {output / 'generation.log'}", file=sys.stderr)
        return 1
    finally:
        handler.close()
        logger.removeHandler(handler)


if __name__ == "__main__":
    raise SystemExit(main())
